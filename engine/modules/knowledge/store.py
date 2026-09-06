"""Local, dependency-light enterprise knowledge-base domain and retrieval service."""
from __future__ import annotations

import csv, hashlib, io, json, math, re, sqlite3, time, uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from ..memory.embedding import HashingEmbeddingModel
from .chunking import chunk_text
from .models import KnowledgeBase, KnowledgeChunk, KnowledgeDocument

ALLOWED_EXTENSIONS={".txt",".md",".markdown",".html",".htm",".pdf",".docx",".xlsx",".csv",".json"}
MAX_FILE_SIZE=20*1024*1024
def now()->str:return datetime.now(timezone.utc).isoformat()
def ident(prefix:str)->str:return f"{prefix}-{uuid.uuid4().hex[:16]}"
def clean_labels(value:Any)->list[str]:return sorted({str(x).strip() for x in (value or []) if str(x).strip()})
def cosine(a:list[float],b:list[float])->float:
    if not a or not b:return 0.0
    return sum(x*y for x,y in zip(a,b))/(math.sqrt(sum(x*x for x in a))*math.sqrt(sum(y*y for y in b)) or 1)
def tokens(text:str)->set[str]:return set(re.findall(r"[\w\u4e00-\u9fff]+",text.lower()))

class KnowledgeStore:
    """SQLite adapter; replaceable behind the same retrieval contract for pgvector/Qdrant."""
    def __init__(self, root:str|Path="runs/knowledge"):
        self.root=Path(root);self.root.mkdir(parents=True,exist_ok=True);(self.root/"files").mkdir(exist_ok=True)
        self.db=self.root/"knowledge.sqlite3";self.conn=sqlite3.connect(self.db,check_same_thread=False);self.conn.row_factory=sqlite3.Row
        self.conn.execute("PRAGMA journal_mode=WAL");self.embedder=HashingEmbeddingModel();self._schema()
    def _schema(self):
        with self.conn:self.conn.executescript("""
        CREATE TABLE IF NOT EXISTS knowledge_bases(id TEXT PRIMARY KEY,data TEXT NOT NULL,owner TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS knowledge_documents(id TEXT PRIMARY KEY,kb_id TEXT NOT NULL,data TEXT NOT NULL,checksum TEXT NOT NULL,UNIQUE(kb_id,checksum));
        CREATE TABLE IF NOT EXISTS knowledge_chunks(id TEXT PRIMARY KEY,kb_id TEXT NOT NULL,document_id TEXT NOT NULL,data TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS knowledge_logs(id TEXT PRIMARY KEY,kb_id TEXT NOT NULL,data TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS idx_kd_kb ON knowledge_documents(kb_id); CREATE INDEX IF NOT EXISTS idx_kc_kb ON knowledge_chunks(kb_id); CREATE INDEX IF NOT EXISTS idx_kl_kb ON knowledge_logs(kb_id);
        """)
    def _row(self,row,kind):return globals()[kind](**json.loads(row["data"]))
    def create_base(self,data:dict[str,Any],owner:str)->KnowledgeBase:
        name=str(data.get("name") or "").strip()
        if not name:raise ValueError("知识库名称不能为空")
        if len(name)>80:raise ValueError("知识库名称不能超过 80 个字符")
        kb=KnowledgeBase(id=ident("kb"),name=name,description=str(data.get("description") or ""),owner_user_id=owner,workspace_id=str(data.get("workspace_id") or "local"),type=str(data.get("type") or "document"),edition=str(data.get("edition") or "standard"),embedding_model=str(data.get("embedding_model") or "hashing"),retrieval_mode=self._mode(data.get("retrieval_mode")),chunk_strategy=str(data.get("chunk_strategy") or "smart"),chunk_size=self._num(data.get("chunk_size"),600,100,4000),chunk_overlap=self._num(data.get("chunk_overlap"),80,0,1000),similarity_threshold=float(data.get("similarity_threshold",.15)),top_k=self._num(data.get("top_k"),5,1,50),rerank_enabled=bool(data.get("rerank_enabled",False)),metadata=dict(data.get("metadata") or {}))
        self._save_base(kb);return kb
    def _save_base(self,kb):
        kb.updated_at=now()
        with self.conn:self.conn.execute("INSERT OR REPLACE INTO knowledge_bases VALUES(?,?,?)",(kb.id,json.dumps(kb.to_dict(),ensure_ascii=False),kb.owner_user_id))
    def list_bases(self,owner:str)->list[KnowledgeBase]:return [self._row(r,"KnowledgeBase") for r in self.conn.execute("SELECT data FROM knowledge_bases WHERE owner=? ORDER BY rowid DESC",(owner,))]
    def get_base(self,kb_id:str,owner:str|None=None)->KnowledgeBase:
        row=self.conn.execute("SELECT data,owner FROM knowledge_bases WHERE id=?",(kb_id,)).fetchone()
        if not row or (owner is not None and row["owner"]!=owner):raise KeyError("知识库不存在")
        return self._row(row,"KnowledgeBase")
    def update_base(self,kb_id:str,data:dict[str,Any],owner:str)->KnowledgeBase:
        kb=self.get_base(kb_id,owner)
        for key in ("name","description","type","edition","embedding_model","retrieval_mode","chunk_strategy","rerank_enabled","metadata"):
            if key in data:setattr(kb,key,self._mode(data[key]) if key=="retrieval_mode" else data[key])
        for key,default,low,high in (("chunk_size",600,100,4000),("chunk_overlap",80,0,1000),("top_k",5,1,50)):
            if key in data:setattr(kb,key,self._num(data[key],default,low,high))
        if "similarity_threshold" in data:kb.similarity_threshold=max(0,min(1,float(data["similarity_threshold"])))
        self._save_base(kb);return kb
    def delete_base(self,kb_id:str,owner:str):
        self.get_base(kb_id,owner)
        for doc in self.list_documents(kb_id):
            path=self.root/doc.source_uri
            if path.exists():path.unlink()
        with self.conn:self.conn.execute("DELETE FROM knowledge_chunks WHERE kb_id=?",(kb_id,));self.conn.execute("DELETE FROM knowledge_documents WHERE kb_id=?",(kb_id,));self.conn.execute("DELETE FROM knowledge_logs WHERE kb_id=?",(kb_id,));self.conn.execute("DELETE FROM knowledge_bases WHERE id=?",(kb_id,))
    def add_document(self,kb_id:str,owner:str,filename:str,raw:bytes,labels:list[str]|None=None)->KnowledgeDocument:
        self.get_base(kb_id,owner);safe=Path(filename).name
        if not safe or safe!=filename or Path(safe).suffix.lower() not in ALLOWED_EXTENSIONS:raise ValueError("不支持的文件类型")
        if len(raw)>MAX_FILE_SIZE:raise ValueError("文件超过 20MB 限制")
        checksum=hashlib.sha256(raw).hexdigest();doc=KnowledgeDocument(id=ident("doc"),knowledge_base_id=kb_id,filename=safe,mime_type=_mime(safe),file_size=len(raw),checksum=checksum,labels=clean_labels(labels))
        try:
            with self.conn:self.conn.execute("INSERT INTO knowledge_documents VALUES(?,?,?,?)",(doc.id,kb_id,json.dumps(doc.to_dict(),ensure_ascii=False),checksum))
        except sqlite3.IntegrityError:raise FileExistsError("相同文件已存在于该知识库")
        rel=f"files/{doc.id}{Path(safe).suffix.lower()}";(self.root/rel).write_bytes(raw);doc.source_uri=rel;self._save_document(doc);self.reparse(kb_id,doc.id,owner);return self.get_document(kb_id,doc.id)
    def _save_document(self,doc):
        doc.updated_at=now()
        with self.conn:self.conn.execute("UPDATE knowledge_documents SET data=? WHERE id=?",(json.dumps(doc.to_dict(),ensure_ascii=False),doc.id))
    def list_documents(self,kb_id:str)->list[KnowledgeDocument]:return [self._row(r,"KnowledgeDocument") for r in self.conn.execute("SELECT data FROM knowledge_documents WHERE kb_id=? ORDER BY rowid DESC",(kb_id,))]
    def get_document(self,kb_id:str,doc_id:str)->KnowledgeDocument:
        row=self.conn.execute("SELECT data FROM knowledge_documents WHERE id=? AND kb_id=?",(doc_id,kb_id)).fetchone()
        if not row:raise KeyError("文档不存在")
        return self._row(row,"KnowledgeDocument")
    def delete_document(self,kb_id:str,doc_id:str,owner:str):
        self.get_base(kb_id,owner);doc=self.get_document(kb_id,doc_id);path=self.root/doc.source_uri
        if path.exists():path.unlink()
        with self.conn:self.conn.execute("DELETE FROM knowledge_chunks WHERE document_id=?",(doc_id,));self.conn.execute("DELETE FROM knowledge_documents WHERE id=?",(doc_id,))
        self._counts(kb_id)
    def reparse(self,kb_id:str,doc_id:str,owner:str):
        self.get_base(kb_id,owner);doc=self.get_document(kb_id,doc_id);doc.parse_status="processing";self._save_document(doc)
        try:
            text,pages=self._extract(self.root/doc.source_uri,doc.filename);doc.parse_status="completed";doc.index_status="processing";self._save_document(doc)
            self._replace_chunks(doc,text,pages);doc.index_status="completed";doc.error_message=""
        except Exception as exc:doc.parse_status="failed";doc.index_status="failed";doc.error_message=str(exc)[:500]
        self._save_document(doc);self._counts(kb_id);return doc
    def reindex(self,kb_id:str,doc_id:str,owner:str):return self.reparse(kb_id,doc_id,owner)
    def _extract(self,path:Path,name:str)->tuple[str,list[int|None]]:
        ext=path.suffix.lower();raw=path.read_bytes()
        if ext in {".txt",".md",".markdown"}:return raw.decode("utf-8",errors="replace"),[]
        if ext in {".html",".htm"}:
            from bs4 import BeautifulSoup
            soup=BeautifulSoup(raw.decode("utf-8",errors="replace"),"html.parser")
            for tag in soup(["script","style","noscript"]):tag.decompose()
            return "\n".join(soup.stripped_strings),[]
        if ext==".csv":return "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(raw.decode("utf-8-sig",errors="replace")))),[]
        if ext==".json":return json.dumps(json.loads(raw.decode("utf-8")),ensure_ascii=False,indent=2),[]
        if ext==".pdf":
            # PyMuPDF preserves CJK character maps for many TeX/embedded-font
            # PDFs where pypdf returns mojibake.  pypdf remains a lightweight
            # fallback for installations that do not include PyMuPDF.
            try:
                import pymupdf
                pdf=pymupdf.open(stream=raw,filetype="pdf")
                pages=[page.get_text("text").strip() for page in pdf]
            except ImportError:
                from pypdf import PdfReader
                pages=[page.extract_text() or "" for page in PdfReader(io.BytesIO(raw)).pages]
            text=self._clean_pdf_text("\n\f\n".join(pages))
            if not self._usable_pdf_text(text):
                raise ValueError("PDF 未提取到可用文本（可能是扫描件或缺少字体映射），请先进行 OCR 后再上传")
            return text,list(range(1,len(pages)+1))
        if ext==".docx":
            from docx import Document
            return "\n".join(p.text for p in Document(io.BytesIO(raw)).paragraphs),[]
        if ext==".xlsx":
            from openpyxl import load_workbook
            book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True);return "\n".join(f"[{s.title}]\n"+"\n".join(" | ".join(str(x or "") for x in row) for row in s.iter_rows(values_only=True)) for s in book.worksheets),[]
        raise ValueError("不支持的文件类型")
    def _replace_chunks(self,doc:KnowledgeDocument,text:str,pages:list[int|None]):
        kb=self.get_base(doc.knowledge_base_id);parts=chunk_text(text,kb.chunk_strategy,kb.chunk_size,kb.chunk_overlap)
        with self.conn:self.conn.execute("DELETE FROM knowledge_chunks WHERE document_id=?",(doc.id,))
        for i,content in enumerate(parts):
            chunk=KnowledgeChunk(id=hashlib.sha256(f"{doc.id}:{i}:{content}".encode()).hexdigest()[:32],knowledge_base_id=doc.knowledge_base_id,document_id=doc.id,content=content,title=doc.filename,page_number=(text[:text.find(content)].count("\f")+1 if "\f" in text else None),chunk_index=i,token_count=len(tokens(content)),embedding=self.embedder.embed(content),labels=doc.labels,metadata={"source_uri":doc.source_uri})
            with self.conn:self.conn.execute("INSERT INTO knowledge_chunks VALUES(?,?,?,?)",(chunk.id,chunk.knowledge_base_id,chunk.document_id,json.dumps(chunk.to_dict(True),ensure_ascii=False)))
    def list_chunks(self,kb_id:str,doc_id:str|None=None)->list[KnowledgeChunk]:
        sql="SELECT data FROM knowledge_chunks WHERE kb_id=?";args=[kb_id]
        if doc_id:sql+=" AND document_id=?";args.append(doc_id)
        return [self._row(r,"KnowledgeChunk") for r in self.conn.execute(sql,args)]
    def save_chunk(self,kb_id:str,data:dict[str,Any],owner:str,chunk_id:str|None=None)->KnowledgeChunk:
        self.get_base(kb_id,owner)
        if chunk_id:
            existing=next((x for x in self.list_chunks(kb_id) if x.id==chunk_id),None)
            if not existing:raise KeyError("切片不存在")
            for k in ("content","title","page_number","labels","metadata","enabled"):
                if k in data:setattr(existing,k,data[k])
            existing.token_count=len(tokens(existing.content));existing.embedding=self.embedder.embed(existing.content);existing.updated_at=now();chunk=existing
        else:
            doc_id=str(data.get("document_id") or "");self.get_document(kb_id,doc_id);content=str(data.get("content") or "").strip()
            if not content:raise ValueError("切片内容不能为空")
            chunk=KnowledgeChunk(id=ident("chunk"),knowledge_base_id=kb_id,document_id=doc_id,content=content,title=str(data.get("title") or ""),chunk_index=int(data.get("chunk_index") or 0),token_count=len(tokens(content)),embedding=self.embedder.embed(content),labels=clean_labels(data.get("labels")),metadata=dict(data.get("metadata") or {}))
        with self.conn:self.conn.execute("INSERT OR REPLACE INTO knowledge_chunks VALUES(?,?,?,?)",(chunk.id,kb_id,chunk.document_id,json.dumps(chunk.to_dict(True),ensure_ascii=False)))
        self._counts(kb_id);return chunk
    def delete_chunk(self,kb_id:str,chunk_id:str,owner:str):
        self.get_base(kb_id,owner)
        with self.conn:self.conn.execute("DELETE FROM knowledge_chunks WHERE id=? AND kb_id=?",(chunk_id,kb_id))
        self._counts(kb_id)
    def retrieve(self,kb_ids:list[str],query:str,*,owner:str,mode:str="hybrid",top_k:int=5,threshold:float=0.15,labels:list[str]|None=None,document_ids:list[str]|None=None,bindings:dict[str,dict]|None=None,application_id:str="",workflow_id:str="",run_id:str="")->dict[str,Any]:
        from ..live_events import emit
        names = [self.get_base(kb_id, owner).name for kb_id in kb_ids]
        emit("knowledge_retrieval_start", knowledge_base_ids=kb_ids, knowledge_base_names=names, message="正在检索知识库：" + "、".join(names))
        started=time.perf_counter();query=" ".join(str(query).split())
        if not query:raise ValueError("query 不能为空")
        allowed=[]
        for kb_id in kb_ids:self.get_base(kb_id,owner);allowed.extend(self.list_chunks(kb_id))
        qemb=self.embedder.embed(query);qt=tokens(query);results=[]
        for c in allowed:
            if not c.enabled or (labels and not set(labels).intersection(c.labels)) or (document_ids and c.document_id not in document_ids):continue
            dense=(cosine(qemb,c.embedding)+1)/2;sparse=len(qt&tokens(c.content))/max(1,len(qt));score=dense if mode=="dense" else sparse if mode=="sparse" else .65*dense+.35*sparse
            weight=float((bindings or {}).get(c.knowledge_base_id,{}).get("weight",1));score*=max(0,weight)
            if score>=threshold:
                d=self.get_document(c.knowledge_base_id,c.document_id);results.append({"chunk_id":c.id,"document_id":c.document_id,"knowledge_base_id":c.knowledge_base_id,"content":c.content,"filename":d.filename,"title":c.title or d.filename,"page_number":c.page_number,"score":round(score,5),"labels":c.labels,"metadata":c.metadata})
        results.sort(key=lambda x:x["score"],reverse=True);results=results[:max(1,min(50,top_k))];latency=round((time.perf_counter()-started)*1000,2)
        log={"id":ident("kr"),"knowledge_base_ids":kb_ids,"application_id":application_id,"workflow_id":workflow_id,"run_id":run_id,"query":query,"rewritten_query":query,"retrieval_mode":mode,"top_k":top_k,"threshold":threshold,"result_count":len(results),"latency_ms":latency,"results":results,"created_at":now()}
        with self.conn:self.conn.execute("INSERT INTO knowledge_logs VALUES(?,?,?)",(log["id"],",".join(kb_ids),json.dumps(log,ensure_ascii=False)))
        emit("knowledge_retrieval_end", knowledge_base_ids=kb_ids, result_count=len(results), duration_ms=latency, message=f"知识库检索完成，命中 {len(results)} 条，耗时 {latency} ms")
        return {"documents":results,"citations":[{k:r[k] for k in ("chunk_id","document_id","knowledge_base_id","filename","title","page_number","score")} for r in results],"context":"\n\n".join(f"[来源 {i+1}: {r['filename']}]\n{r['content']}" for i,r in enumerate(results)),"retrieval_metadata":{"query":query,"mode":mode,"latency_ms":latency,"result_count":len(results),"log_id":log["id"]}}
    def logs(self,kb_id:str)->list[dict[str,Any]]:return [json.loads(r["data"]) for r in self.conn.execute("SELECT data FROM knowledge_logs WHERE kb_id LIKE ? ORDER BY rowid DESC",(f"%{kb_id}%",))]
    def statistics(self,kb_id:str)->dict[str,Any]:
        kb=self.get_base(kb_id);return {"knowledge_base_id":kb_id,"document_count":kb.document_count,"chunk_count":kb.chunk_count,"status":kb.status,"parse_failed_count":sum(d.parse_status=="failed" for d in self.list_documents(kb_id)),"retrieval_count":len(self.logs(kb_id)),"updated_at":kb.updated_at}
    def _counts(self,kb_id):
        kb=self.get_base(kb_id);kb.document_count=len(self.list_documents(kb_id));kb.chunk_count=len(self.list_chunks(kb_id));kb.status="ready";self._save_base(kb)
    @staticmethod
    def _mode(v):return str(v) if str(v) in {"dense","sparse","hybrid"} else "hybrid"
    @staticmethod
    def _usable_pdf_text(text:str)->bool:
        """Reject empty/obviously-corrupted extraction instead of indexing it."""
        visible=re.findall(r"[A-Za-z0-9\u4e00-\u9fff]",text)
        if len(visible)<12:return False
        replacement=text.count("\ufffd")
        return replacement/max(1,len(text))<0.08
    @staticmethod
    def _clean_pdf_text(text:str)->str:
        """Drop PDF layout artefacts such as table-of-contents dot leaders."""
        pages=[]
        for page in text.split("\f"):
            cleaned=[]
            for line in page.splitlines():
                compact=line.strip()
                # A text extractor may emit ``. . . .`` as one line or one
                # dot per line.  Neither form is meaningful knowledge.
                compact=re.sub(r"(?:[.·…]\s*){3,}"," ",compact)
                compact=re.sub(r"\s{2,}"," ",compact).strip()
                if not compact or re.fullmatch(r"[.。·•…\-_\s]+",compact):
                    continue
                cleaned.append(compact)
            pages.append("\n".join(cleaned))
        return "\f".join(pages)
    @staticmethod
    def _num(v,d,low,high):
        try:return max(low,min(high,int(v)))
        except (TypeError,ValueError):return d

def _mime(name:str)->str:return {".txt":"text/plain",".md":"text/markdown",".html":"text/html",".csv":"text/csv",".json":"application/json",".pdf":"application/pdf",".docx":"application/vnd.openxmlformats-officedocument.wordprocessingml.document",".xlsx":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}.get(Path(name).suffix.lower(),"application/octet-stream")

