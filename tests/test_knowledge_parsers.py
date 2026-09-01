import io
import json

from docx import Document
from openpyxl import Workbook
import pymupdf

from engine.modules.knowledge import KnowledgeStore


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("员工手册", 1)
    document.add_paragraph("年假政策为每年十天带薪年假。")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _xlsx_bytes() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "人员"
    sheet.append(["姓名", "部门"])
    sheet.append(["张三", "研发部"])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _pdf_bytes() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Knowledge base PDF extraction verifies searchable text.")
    return document.tobytes()


def test_supported_document_formats_are_parsed_and_indexed(tmp_path):
    store = KnowledgeStore(tmp_path / "knowledge")
    base = store.create_base({"name": "解析回归", "chunk_size": 100}, "user-a")
    inputs = {
        "note.txt": "纯文本资料可以检索。".encode(),
        "guide.md": "# Markdown 标题\n\nMarkdown 正文可以检索。".encode(),
        "page.html": "<h1>HTML 标题</h1><script>ignored()</script><p>HTML 正文可以检索。</p>".encode(),
        "table.csv": "姓名,部门\n张三,研发部\n".encode("utf-8"),
        "record.json": json.dumps({"标题": "JSON 资料", "内容": "JSON 正文可以检索"}, ensure_ascii=False).encode(),
        "manual.docx": _docx_bytes(),
        "members.xlsx": _xlsx_bytes(),
        "manual.pdf": _pdf_bytes(),
    }
    for filename, raw in inputs.items():
        document = store.add_document(base.id, "user-a", filename, raw)
        assert document.parse_status == "completed", document.error_message
        assert document.index_status == "completed"
        assert store.list_chunks(base.id, document.id)

    html_doc = next(item for item in store.list_documents(base.id) if item.filename == "page.html")
    html_text = "\n".join(chunk.content for chunk in store.list_chunks(base.id, html_doc.id))
    assert "HTML 正文" in html_text
    assert "ignored" not in html_text


def test_pdf_layout_leader_dots_do_not_become_chunks(tmp_path):
    store = KnowledgeStore(tmp_path / "knowledge")
    base = store.create_base({"name": "PDF 清洗", "chunk_size": 100}, "user-a")
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Contents\nIntroduction\n.\n.\n.\n.\nChapter one searchable content")
    uploaded = store.add_document(base.id, "user-a", "toc.pdf", document.tobytes())
    contents = "\n".join(item.content for item in store.list_chunks(base.id, uploaded.id))
    assert "Chapter one searchable content" in contents
    assert "\n.\n" not in contents


def test_pdf_cleaner_removes_inline_table_of_contents_leaders(tmp_path):
    store = KnowledgeStore(tmp_path / "knowledge")
    cleaned = store._clean_pdf_text("目录\n第一章 . . . . . . 12\n第二章\f正文")
    assert cleaned == "目录\n第一章 12\n第二章\f正文"
